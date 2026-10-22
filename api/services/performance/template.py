"""The RoomRate bookings template (Excel): what the owner fills, and what the importer reads.

Four sheets: «Κρατήσεις» (one row per booking), «Δωμάτια» (rooms per room
type, prefilled with the property's room types), «Κλειστές περίοδοι» and
«Οδηγίες». The column lists below are the one definition of the template:
the builder writes them and api/services/performance/importing reads them, by header.

Rows of «Κρατήσεις» whose booking number starts with «ΠΑΡΑΔΕΙΓΜΑ» are the
grey examples: the importer skips them, so a template sent back with them
still in place imports cleanly.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

BOOKINGS_SHEET = "Κρατήσεις"
ROOMS_SHEET = "Δωμάτια"
CLOSURES_SHEET = "Κλειστές περίοδοι"
INSTRUCTIONS_SHEET = "Οδηγίες"
TEMPLATE_FILENAME = "RoomRate_kratiseis.xlsx"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXAMPLE_PREFIX = "ΠΑΡΑΔΕΙΓΜΑ"
MAX_ROWS = 20_000

# What the owner picks in «Κατάσταση» → the stored status.
STATUSES = {"Επιβεβαιωμένη": "confirmed", "Ακυρωμένη": "cancelled", "No-show": "no_show"}
STATUS_LABELS = {code: label for label, code in STATUSES.items()}
# «Κανάλι» suggestions; any other text is kept as typed.
CHANNELS = ("Booking.com", "Expedia", "Airbnb", "Απευθείας", "Τηλέφωνο / email", "Πρακτορείο", "Άλλο")
DEFAULT_CHANNEL = "Άλλο"

_SEA = "0B5D7A"
_GREY = "8A8F98"
_DATE_FORMAT = "DD/MM/YYYY"
_MONEY_FORMAT = "#,##0.00"


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    kind: str  # text | date | int | money | status | channel | room
    required: bool
    width: int
    help: str


BOOKING_COLUMNS: tuple[Column, ...] = (
    Column("booking_ref", "Αριθμός κράτησης", "text", True, 18,
           "Ο αριθμός της κράτησης (από το κανάλι ή δικός σας). "
           "Ίδιος αριθμός σε νέο ανέβασμα = ενημέρωση, όχι διπλή κράτηση."),
    Column("booked_on", "Ημερομηνία κράτησης", "date", False, 14,
           "Πότε έγινε η κράτηση. Χωρίς αυτήν δεν υπολογίζεται ο χρόνος κράτησης."),
    Column("arrival", "Άφιξη", "date", True, 12, "Ημερομηνία άφιξης."),
    Column("departure", "Αναχώρηση", "date", True, 12,
           "Ημερομηνία αναχώρησης (η τελευταία νύχτα είναι η προηγούμενη μέρα)."),
    Column("room_label", "Τύπος δωματίου", "room", True, 30,
           "Από τη λίστα (φύλλο «Δωμάτια») ή όπως τον γράφει το κανάλι· τον αντιστοιχίζετε μία φορά στο ανέβασμα."),
    Column("rooms", "Δωμάτια", "int", False, 9, "Πόσα δωμάτια αυτού του τύπου (κενό = 1)."),
    Column("adults", "Ενήλικες", "int", False, 9, "Προαιρετικό."),
    Column("children", "Παιδιά", "int", False, 8, "Προαιρετικό."),
    Column("gross_eur", "Συνολικό ποσό (€)", "money", True, 16,
           "Ό,τι πλήρωσε ο πελάτης για όλη τη διαμονή και όλα τα δωμάτια, με ΦΠΑ."),
    Column("commission_eur", "Προμήθεια καναλιού (€)", "money", False, 18,
           "Η προμήθεια του καναλιού για την κράτηση (κενό = 0)."),
    Column("channel", "Κανάλι", "channel", False, 18, "Από τη λίστα ή άλλο κείμενο (κενό = «Άλλο»)."),
    Column("status", "Κατάσταση", "status", False, 15, "Επιβεβαιωμένη, Ακυρωμένη ή No-show (κενό = Επιβεβαιωμένη)."),
    Column("cancelled_on", "Ημερομηνία ακύρωσης", "date", False, 14, "Μόνο για ακυρωμένες."),
    Column("country", "Χώρα πελάτη", "text", False, 18,
           "Ελληνικά ή αγγλικά ή κωδικός (GR, DE). Ποτέ όνομα, email ή τηλέφωνο πελάτη."),
)

ROOM_COLUMNS: tuple[Column, ...] = (
    Column("room_label", "Τύπος δωματίου", "text", True, 34, "Όπως στο φύλλο «Κρατήσεις»."),
    Column("rooms", "Αριθμός δωματίων", "int", True, 17, "Πόσα δωμάτια έχει ο τύπος."),
    Column("valid_from", "Ισχύει από", "date", False, 13, "Μόνο αν άλλαξε μέσα στη χρονιά (κενό = πάντα)."),
    Column("valid_to", "Ισχύει έως", "date", False, 13, "Μόνο αν άλλαξε μέσα στη χρονιά (κενό = πάντα)."),
)

CLOSURE_COLUMNS: tuple[Column, ...] = (
    Column("date_from", "Από", "date", True, 12, "Πρώτη μέρα που ήταν κλειστό."),
    Column("date_to", "Έως", "date", True, 12, "Τελευταία μέρα που ήταν κλειστό."),
    Column("room_label", "Τύπος δωματίου", "room", False, 30, "Κενό = όλο το κατάλυμα."),
    Column("rooms", "Δωμάτια", "int", False, 9, "Κενό = όλα τα δωμάτια (του τύπου ή του καταλύματος)."),
    Column("reason", "Αιτία", "text", False, 28, "Προαιρετικό, π.χ. «Εκτός σεζόν»."),
)

_EXAMPLES: tuple[dict[str, Any], ...] = (
    {"booking_ref": f"{EXAMPLE_PREFIX} 1", "booked_on": date(2026, 3, 2), "arrival": date(2026, 7, 10),
     "departure": date(2026, 7, 14), "rooms": 1, "adults": 2, "children": 0, "gross_eur": 520.0,
     "commission_eur": 88.4, "channel": "Booking.com", "status": "confirmed", "country": "Γερμανία"},
    {"booking_ref": f"{EXAMPLE_PREFIX} 2", "booked_on": date(2026, 5, 20), "arrival": date(2026, 8, 3),
     "departure": date(2026, 8, 10), "rooms": 2, "adults": 4, "children": 2, "gross_eur": 1890.0,
     "commission_eur": 0, "channel": "Απευθείας", "status": "confirmed", "country": "Ελλάδα"},
    {"booking_ref": f"{EXAMPLE_PREFIX} 3", "booked_on": date(2026, 4, 11), "arrival": date(2026, 6, 5),
     "departure": date(2026, 6, 8), "rooms": 1, "adults": 2, "children": 0, "gross_eur": 330.0,
     "commission_eur": 56.1, "channel": "Booking.com", "status": "cancelled", "cancelled_on": date(2026, 5, 28),
     "country": "UK"},
)


def build_template(
    room_types: Sequence[dict[str, Any]],
    *,
    property_name: str | None = None,
    bookings: Iterable[dict[str, Any]] | None = None,
    closures: Iterable[dict[str, Any]] = (),
) -> bytes:
    """The template as .xlsx bytes.

    ``room_types``: ``{"label", "rooms"}`` (optionally ``valid_from`` /
    ``valid_to``) per room type of the property, rooms None when not known,
    prefilled in «Δωμάτια» and offered in the room lists. ``bookings`` None writes the grey examples; given (the demo
    generator), they are written instead, ``status`` as a code.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = BOOKINGS_SHEET
    _header(sheet, BOOKING_COLUMNS)
    if bookings is None:
        labels = [room["label"] for room in room_types] or ["Δίκλινο"]
        rows = [{**example, "room_label": labels[i % len(labels)]} for i, example in enumerate(_EXAMPLES)]
        _write_rows(sheet, BOOKING_COLUMNS, rows, font=Font(color=_GREY, italic=True))
    else:
        _write_rows(sheet, BOOKING_COLUMNS, bookings)
    _validate(sheet, BOOKING_COLUMNS)

    rooms = workbook.create_sheet(ROOMS_SHEET)
    _header(rooms, ROOM_COLUMNS)
    _write_rows(rooms, ROOM_COLUMNS, ({**room, "room_label": room["label"]} for room in room_types))
    _validate(rooms, ROOM_COLUMNS)

    closed = workbook.create_sheet(CLOSURES_SHEET)
    _header(closed, CLOSURE_COLUMNS)
    _write_rows(closed, CLOSURE_COLUMNS, closures)
    _validate(closed, CLOSURE_COLUMNS)

    _instructions(workbook.create_sheet(INSTRUCTIONS_SHEET), property_name)
    workbook.active = 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _header(sheet: Worksheet, columns: Sequence[Column]) -> None:
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column.header + (" *" if column.required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_SEA)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"


def _write_rows(
    sheet: Worksheet,
    columns: Sequence[Column],
    rows: Iterable[dict[str, Any]],
    *,
    font: Font | None = None,
) -> None:
    for row_number, row in enumerate(rows, start=2):
        for index, column in enumerate(columns, start=1):
            value = row.get(column.key)
            if column.kind == "status" and value is not None:
                value = STATUS_LABELS.get(value, value)
            cell = sheet.cell(row=row_number, column=index, value=value)
            if column.kind == "date":
                cell.number_format = _DATE_FORMAT
            elif column.kind == "money":
                cell.number_format = _MONEY_FORMAT
            if font is not None:
                cell.font = font


def _validate(sheet: Worksheet, columns: Sequence[Column]) -> None:
    """Lists and checks on every fillable row; a channel or room outside the list is allowed (a warning)."""
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        cells = f"{letter}2:{letter}{MAX_ROWS + 1}"
        if column.kind == "status":
            rule = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
            rule.error, rule.errorTitle = "Διαλέξτε από τη λίστα.", "Κατάσταση"
        elif column.kind == "channel":
            rule = DataValidation(type="list", formula1='"' + ",".join(CHANNELS) + '"', allow_blank=True)
            rule.errorStyle, rule.error, rule.errorTitle = "warning", "Δεν είναι στη λίστα· να μείνει;", "Κανάλι"
        elif column.kind == "room":
            rule = DataValidation(type="list", formula1=f"'{ROOMS_SHEET}'!$A$2:$A$200", allow_blank=True)
            rule.errorStyle, rule.error, rule.errorTitle = (
                "warning", "Δεν είναι στο φύλλο «Δωμάτια»· θα τον αντιστοιχίσετε στο ανέβασμα.", "Τύπος δωματίου")
        elif column.kind == "date":
            # 36526 = 1/1/2000 as an Excel date.
            rule = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
            rule.errorStyle, rule.error, rule.errorTitle = "warning", "Γράψτε ημερομηνία, π.χ. 14/07/2026.", "Ημερομηνία"
        elif column.kind in {"int", "money"}:
            rule = DataValidation(
                type="whole" if column.kind == "int" else "decimal",
                operator="greaterThanOrEqual",
                formula1="0",
                allow_blank=True,
            )
            rule.error, rule.errorTitle = "Γράψτε αριθμό (0 ή μεγαλύτερο).", column.header
        else:
            continue
        rule.showErrorMessage = True
        rule.add(cells)
        sheet.add_data_validation(rule)


def _instructions(sheet: Worksheet, property_name: str | None) -> None:
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 90
    title = "Πρότυπο κρατήσεων RoomRate" + (f" · {property_name}" if property_name else "")
    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color=_SEA)
    steps = (
        "1. Στο φύλλο «Κρατήσεις» μία γραμμή ανά κράτηση: όλες, και οι ακυρωμένες (μετράνε στο ποσοστό ακυρώσεων).",
        "2. Στο φύλλο «Δωμάτια» πόσα δωμάτια έχει κάθε τύπος: χωρίς αυτά δεν υπολογίζεται πληρότητα.",
        "3. Στο φύλλο «Κλειστές περίοδοι» πότε ήταν κλειστό (π.χ. χειμώνας): αυτές οι μέρες δεν μετράνε ως απούλητες.",
        "4. Οι γκρι γραμμές «ΠΑΡΑΔΕΙΓΜΑ» αγνοούνται· σβήστε τες ή αφήστε τες.",
        "5. Ανεβάστε το αρχείο στο RoomRate → Απόδοση → Ανέβασμα κρατήσεων. Ξανανέβασμα = ενημέρωση, ποτέ διπλές κρατήσεις.",
        "Τα στοιχεία των πελατών (όνομα, email, τηλέφωνο, διεύθυνση) δεν χρειάζονται και δεν διαβάζονται ποτέ.",
    )
    for offset, step in enumerate(steps, start=3):
        sheet.cell(row=offset, column=1, value=step)
    row = len(steps) + 4
    for index, header in enumerate(("Στήλη (Κρατήσεις)", "Υποχρεωτική", "Τι γράφω"), start=1):
        cell = sheet.cell(row=row, column=index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_SEA)
    for column in BOOKING_COLUMNS:
        row += 1
        sheet.cell(row=row, column=1, value=column.header)
        sheet.cell(row=row, column=2, value="Ναι" if column.required else "")
        sheet.cell(row=row, column=3, value=column.help).alignment = Alignment(wrap_text=True)

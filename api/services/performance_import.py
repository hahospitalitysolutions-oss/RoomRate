"""Read an owner's bookings file: the RoomRate template or the Booking.com reservations export.

Both formats, as .xlsx or .csv, end as the same bookings (the fields of
PerformanceRepository.save_bookings without ``room_type_id``, which the
owner's label mapping supplies) plus one Greek message per row that cannot
be imported, with the row number the owner sees in Excel.

Only the columns listed here are read, by their header: guests' names,
emails, phones and addresses are never looked at. The header row may sit
below a title; a file whose headers are the other format's is read as
that format. Dates come as Excel dates, serial numbers or text
(14/07/2026, 2026-07-14, 14.07.2026, with or without a time); amounts as
numbers or text with either decimal mark and a currency («1.234,50 €»,
«436 EUR»).

Contract: docs/performance/API.md («Imports in detail»).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import load_workbook

from api.services.performance_countries import country_code
from api.services.performance_template import (
    BOOKING_COLUMNS,
    BOOKINGS_SHEET,
    CLOSURE_COLUMNS,
    CLOSURES_SHEET,
    DEFAULT_CHANNEL,
    EXAMPLE_PREFIX,
    MAX_ROWS,
    ROOM_COLUMNS,
    ROOMS_SHEET,
    STATUSES,
)

BOOKING_COM = "Booking.com"
MAX_NIGHTS = 365
_HEADER_SEARCH_ROWS = 10
_EXCEL_EPOCH = date(1899, 12, 30)


class ImportFileError(ValueError):
    """The file as a whole cannot be read; the message is Greek, for the owner (HTTP 400)."""


@dataclass
class ParsedRow:
    row: int
    booking_ref: str | None
    room_label: str | None
    booking: dict[str, Any] | None
    error: str | None = None


@dataclass
class ParsedFile:
    format: str
    rows: list[ParsedRow] = field(default_factory=list)
    # Template only: «Δωμάτια» and «Κλειστές περίοδοι», with labels still unmapped.
    rooms: list[dict[str, Any]] = field(default_factory=list)
    closures: list[dict[str, Any]] = field(default_factory=list)
    # Rows of «Δωμάτια» / «Κλειστές περίοδοι» that cannot be used: {"row", "booking_ref": None, "message"}.
    sheet_errors: list[dict[str, Any]] = field(default_factory=list)

    @property
    def bookings(self) -> list[dict[str, Any]]:
        return [row.booking for row in self.rows if row.booking is not None]

    @property
    def errors(self) -> list[dict[str, Any]]:
        """Every row that will not be imported, bookings first: the preview's ``errors``."""
        return [
            {"row": row.row, "booking_ref": row.booking_ref, "message": row.error}
            for row in self.rows
            if row.error
        ] + self.sheet_errors


# ---------------------------------------------------------------------------
# Columns of each format: field → accepted headers (compared without case, accents or a trailing «*»)
# ---------------------------------------------------------------------------

TEMPLATE_FIELDS: dict[str, tuple[str, ...]] = {column.key: (column.header,) for column in BOOKING_COLUMNS}
TEMPLATE_REQUIRED = tuple(column.key for column in BOOKING_COLUMNS if column.required)

# The extranet export in English, and the names it has had or a Greek extranet may use (task B3: to be
# confirmed against a real export).
BOOKING_FIELDS: dict[str, tuple[str, ...]] = {
    "booking_ref": ("Book number", "Booking number", "Reservation number", "Αριθμός κράτησης"),
    "booked_on": ("Booked on", "Booking date", "Reservation date", "Ημερομηνία κράτησης", "Κρατήθηκε στις"),
    "arrival": ("Check-in", "Arrival", "Άφιξη", "Check-in date"),
    "departure": ("Check-out", "Departure", "Αναχώρηση", "Check-out date"),
    "room_label": ("Unit type", "Room type", "Room name", "Rooms type", "Τύπος δωματίου", "Τύπος μονάδας"),
    "rooms": ("Rooms", "Number of rooms", "Δωμάτια"),
    "adults": ("Adults", "Ενήλικες"),
    "children": ("Children", "Παιδιά"),
    "people": ("People", "Persons", "Guests", "Άτομα"),
    "gross_eur": ("Price", "Total price", "Total payment", "Τιμή", "Συνολική τιμή"),
    "commission_eur": ("Commission amount", "Commission", "Ποσό προμήθειας", "Προμήθεια"),
    "commission_pct": ("Commission %", "Commission percentage", "Προμήθεια %", "Ποσοστό προμήθειας"),
    "status": ("Status", "Κατάσταση"),
    "cancelled_on": ("Cancellation date", "Cancelled on", "Ημερομηνία ακύρωσης"),
    "country": ("Booker country", "Country", "Χώρα", "Χώρα κράτησης"),
    "nights": ("Duration (nights)", "Nights", "Length of stay", "Νύχτες", "Διάρκεια (νύχτες)"),
}
BOOKING_REQUIRED = ("booking_ref", "arrival", "room_label", "gross_eur")

FIELD_NAMES = {
    "booking_ref": "Αριθμός κράτησης",
    "arrival": "Άφιξη",
    "departure": "Αναχώρηση",
    "room_label": "Τύπος δωματίου",
    "gross_eur": "Ποσό",
}

_STATUS_WORDS = {
    **{label: code for label, code in STATUSES.items()},
    "ok": "confirmed",
    "confirmed": "confirmed",
    "επιβεβαιωμενη": "confirmed",
    "επιβεβαιωμενο": "confirmed",
    "cancelled": "cancelled",
    "canceled": "cancelled",
    "ακυρωμενη": "cancelled",
    "ακυρωση": "cancelled",
    "no_show": "no_show",
    "no-show": "no_show",
    "no show": "no_show",
    "noshow": "no_show",
}


def plain_text(value: object) -> str:
    """Lower case, no accents, single spaces, no trailing «*»: how headers and words are compared."""
    text = unicodedata.normalize("NFD", str(value)).strip().rstrip("*").strip()
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.casefold().split())


def _lookup(headers: Sequence[object], fields: dict[str, tuple[str, ...]]) -> dict[str, int]:
    """Field → column index of the first header that names it."""
    plain = [plain_text(header) if header is not None else "" for header in headers]
    found: dict[str, int] = {}
    for key, names in fields.items():
        for name in names:
            if plain_text(name) in plain:
                found[key] = plain.index(plain_text(name))
                break
    return found


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------

_DATE_FORMATS = (
    "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S",
    "%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%y",
    "%d.%m.%Y", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%y",
    "%d-%m-%Y", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
    "%Y/%m/%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
)


class _BadValue(ValueError):
    pass


def _empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_date(value: object) -> date | None:
    """A date from an Excel cell or text; None for an empty cell; _BadValue otherwise."""
    if _empty(value):
        return None
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not 20_000 <= value <= 80_000:  # an Excel serial date between 1954 and 2119
            raise _BadValue(str(value))
        result = _EXCEL_EPOCH + timedelta(days=int(value))
    else:
        text = " ".join(str(value).split())
        for pattern in _DATE_FORMATS:
            try:
                result = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        else:
            raise _BadValue(text)
    if not 2000 <= result.year <= 2100:
        raise _BadValue(str(value))
    return result


_CURRENCY = re.compile(r"(?i)\s*(€|euros?|eur|ευρώ|ευρω)\s*")
_OTHER_CURRENCY = re.compile(r"[A-Za-z]{3}|[$£¥₺]")


def parse_amount(value: object) -> float | None:
    """Euros from a number or text («1.234,50 €», «1,234.50», «436 EUR»); None for an empty cell."""
    if _empty(value):
        return None
    if isinstance(value, bool):
        raise _BadValue(str(value))
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = _CURRENCY.sub("", str(value)).replace(" ", "").replace(" ", "")
    if not re.fullmatch(r"-?[\d.,]+", text) or not re.search(r"\d", text):
        raise _BadValue(str(value).strip())
    if "," in text and "." in text:
        decimal_mark = "," if text.rfind(",") > text.rfind(".") else "."
    elif text.count(",") == 1 or text.count(".") == 1:
        mark = "," if "," in text else "."
        # One mark followed by exactly three digits groups thousands («1.234», «1,234»).
        decimal_mark = None if re.search(r"\d[.,]\d{3}$", text) else mark
    else:
        decimal_mark = None  # «1.234.567» or no mark at all
    thousands = {",", "."} - {decimal_mark}
    for mark in thousands:
        text = text.replace(mark, "")
    if decimal_mark == ",":
        text = text.replace(",", ".")
    return round(float(text), 2)


def parse_int(value: object) -> int | None:
    if _empty(value):
        return None
    if isinstance(value, bool):
        raise _BadValue(str(value))
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, int):
        return value
    text = str(value).strip()
    if re.fullmatch(r"\d+([.,]0+)?", text):
        return int(re.split(r"[.,]", text)[0])
    raise _BadValue(text)


def parse_reference(value: object) -> str | None:
    """A booking number as text: 4018833201 stays «4018833201», never «4018833201.0»."""
    if _empty(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())[:100]


def parse_status(value: object) -> str:
    """confirmed / cancelled / no_show from the template's Greek words or the export's codes."""
    if _empty(value):
        return "confirmed"
    word = plain_text(value)
    for label, code in _STATUS_WORDS.items():
        if plain_text(label) == word:
            return code
    if word.startswith("cancel") or word.startswith("ακυρ"):
        return "cancelled"  # cancelled_by_guest, cancelled_by_hotel, …
    raise _BadValue(str(value).strip())


def _label(value: object) -> str | None:
    if _empty(value):
        return None
    return " ".join(str(value).split())[:255]


def net_eur(
    gross_eur: float,
    commission_eur: float,
    *,
    rooms: int,
    nights: int,
    vat_pct: float,
    resilience_fee_eur: float,
    fee_included_in_price: bool,
) -> float:
    """What the property keeps: gross − resilience fee (when in the price) − VAT − commission."""
    fee = resilience_fee_eur * rooms * nights if fee_included_in_price else 0.0
    return round((gross_eur - fee) / (1 + vat_pct / 100) - commission_eur, 2)


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------


def _sheets(content: bytes, filename: str) -> dict[str, list[tuple]]:
    """Every sheet's rows as tuples (a .csv is one sheet)."""
    name = filename.lower()
    if content[:4] == b"PK\x03\x04":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # a damaged or password-protected workbook
            raise ImportFileError(
                "Το αρχείο Excel δεν διαβάζεται. Ανοίξτε το και αποθηκεύστε το ξανά ως .xlsx."
            ) from exc
        try:
            return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}
        finally:
            workbook.close()
    if content[:4] == b"\xd0\xcf\x11\xe0" or name.endswith(".xls"):
        raise ImportFileError(
            "Το αρχείο είναι παλιού τύπου Excel (.xls). Ανοίξτε το στο Excel και αποθηκεύστε το ως .xlsx ή .csv."
        )
    if name.endswith((".csv", ".txt")):
        return {"csv": _csv_rows(content)}
    raise ImportFileError("Το αρχείο δεν είναι .xlsx ή .csv.")


def _csv_rows(content: bytes) -> list[tuple]:
    for encoding in ("utf-8-sig", "cp1253", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in text:
        raise ImportFileError("Το αρχείο δεν είναι .xlsx ή .csv.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    return [tuple(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _find_header(rows: Sequence[tuple], fields: dict[str, tuple[str, ...]], required: Sequence[str]):
    """(index of the header row, its columns) among the first rows; columns {} when none has the required."""
    best: tuple[int, dict[str, int]] = (0, {})
    for index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        columns = _lookup(row, fields)
        if all(key in columns for key in required):
            return index, columns
        if len(columns) > len(best[1]):
            best = (index, columns)
    return best[0], {}


def _missing(rows: Sequence[tuple], fields, required) -> list[str]:
    best: dict[str, int] = {}
    for row in rows[:_HEADER_SEARCH_ROWS]:
        columns = _lookup(row, fields)
        if len(columns) > len(best):
            best = columns
    return [FIELD_NAMES.get(key, key) for key in required if key not in best]


def _columns_of(fmt: str) -> tuple[dict[str, tuple[str, ...]], tuple[str, ...]]:
    return (TEMPLATE_FIELDS, TEMPLATE_REQUIRED) if fmt == "template" else (BOOKING_FIELDS, BOOKING_REQUIRED)


def read_bookings_file(content: bytes, filename: str, requested_format: str) -> ParsedFile:
    """Every booking of the file, or ImportFileError when the file itself cannot be used."""
    if not content:
        raise ImportFileError("Το αρχείο είναι άδειο.")
    sheets = _sheets(content, filename)
    formats = ("template", "booking") if requested_format == "template" else ("booking", "template")
    for fmt in formats:
        fields, required = _columns_of(fmt)
        ordered = sorted(sheets.items(), key=lambda item: item[0] != BOOKINGS_SHEET)
        for _title, rows in ordered:
            header_index, columns = _find_header(rows, fields, required)
            if columns:
                parsed = ParsedFile(format=fmt)
                parsed.rows = list(_read_bookings(rows, header_index, columns, fmt))
                if fmt == "template":
                    _read_template_extras(sheets, parsed)
                return parsed
    first = next(iter(sheets.values()), [])
    fields, required = _columns_of(requested_format)
    missing = _missing(first, fields, required)
    hint = (
        "Κατεβάστε το πρότυπο του RoomRate και συμπληρώστε το."
        if requested_format == "template"
        else "Ανεβάστε το export κρατήσεων του extranet όπως κατέβηκε."
    )
    raise ImportFileError(f"Δεν βρέθηκαν οι στήλες: {', '.join(missing)}. {hint}")


def _cell(row: tuple, columns: dict[str, int], key: str) -> object:
    index = columns.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def _read_bookings(rows: Sequence[tuple], header_index: int, columns: dict[str, int], fmt: str) -> Iterator[ParsedRow]:
    data_rows = [(header_index + 2 + offset, row) for offset, row in enumerate(rows[header_index + 1:])]
    relevant = [
        (number, row) for number, row in data_rows
        if any(not _empty(_cell(row, columns, key)) for key in columns)
    ]
    if len(relevant) > MAX_ROWS:
        count, limit = f"{len(relevant):,}".replace(",", "."), f"{MAX_ROWS:,}".replace(",", ".")
        raise ImportFileError(f"Το αρχείο έχει {count} κρατήσεις· το όριο είναι {limit}.")
    seen: dict[str, ParsedRow] = {}
    for number, row in relevant:
        reference = parse_reference(_cell(row, columns, "booking_ref"))
        if fmt == "template" and reference and plain_text(reference).startswith(plain_text(EXAMPLE_PREFIX)):
            continue
        for parsed in _booking_rows(number, row, columns, fmt, reference):
            if parsed.booking is not None and parsed.booking_ref in seen:
                earlier = seen[parsed.booking_ref]
                earlier.booking = None
                earlier.error = f"Ο ίδιος αριθμός κράτησης υπάρχει και στη γραμμή {number}· κρατήθηκε εκείνη."
            if parsed.booking is not None:
                seen[parsed.booking_ref] = parsed
            yield parsed


def _booking_rows(number: int, row: tuple, columns: dict[str, int], fmt: str, reference: str | None) -> list[ParsedRow]:
    label = _label(_cell(row, columns, "room_label"))

    def failed(message: str) -> list[ParsedRow]:
        return [ParsedRow(number, reference, label, None, message)]

    if reference is None:
        return failed("Λείπει ο αριθμός κράτησης.")
    if label is None:
        return failed("Λείπει ο τύπος δωματίου.")
    try:
        arrival = parse_date(_cell(row, columns, "arrival"))
    except _BadValue as exc:
        return failed(f"Η άφιξη δεν είναι ημερομηνία («{exc}»).")
    try:
        departure = parse_date(_cell(row, columns, "departure"))
    except _BadValue as exc:
        return failed(f"Η αναχώρηση δεν είναι ημερομηνία («{exc}»).")
    try:
        nights_given = parse_int(_cell(row, columns, "nights"))
    except _BadValue:
        nights_given = None
    if arrival is None:
        return failed("Λείπει η ημερομηνία άφιξης.")
    if departure is None and nights_given:
        departure = arrival + timedelta(days=nights_given)
    if departure is None:
        return failed("Λείπει η ημερομηνία αναχώρησης.")
    if departure < arrival:
        return failed("Η αναχώρηση είναι πριν από την άφιξη.")
    if departure == arrival:
        return failed("Η αναχώρηση είναι την ίδια μέρα με την άφιξη (καμία νύχτα).")
    nights = (departure - arrival).days
    if nights > MAX_NIGHTS:
        return failed(f"Διαμονή {nights} νυχτών· το όριο είναι {MAX_NIGHTS}.")
    try:
        booked_on = parse_date(_cell(row, columns, "booked_on"))
    except _BadValue as exc:
        return failed(f"Η ημερομηνία κράτησης δεν είναι ημερομηνία («{exc}»).")
    if booked_on is not None and booked_on > arrival:
        return failed("Η ημερομηνία κράτησης είναι μετά την άφιξη.")
    try:
        cancelled_on = parse_date(_cell(row, columns, "cancelled_on"))
    except _BadValue as exc:
        return failed(f"Η ημερομηνία ακύρωσης δεν είναι ημερομηνία («{exc}»).")
    try:
        status = parse_status(_cell(row, columns, "status"))
    except _BadValue as exc:
        return failed(f"Άγνωστη κατάσταση («{exc}»): Επιβεβαιωμένη, Ακυρωμένη ή No-show.")
    raw_gross = _cell(row, columns, "gross_eur")
    if isinstance(raw_gross, str) and _OTHER_CURRENCY.search(_CURRENCY.sub(" ", raw_gross)):
        if re.search(r"\d", raw_gross):
            return failed(f"Το ποσό δεν είναι σε ευρώ («{raw_gross.strip()}»).")
    try:
        gross = parse_amount(raw_gross)
    except _BadValue as exc:
        return failed(f"Το ποσό δεν είναι αριθμός («{exc}»).")
    if gross is None:
        return failed("Λείπει το ποσό.")
    try:
        commission = parse_amount(_cell(row, columns, "commission_eur"))
        commission_pct = parse_amount(_cell(row, columns, "commission_pct"))
    except _BadValue as exc:
        return failed(f"Η προμήθεια δεν είναι αριθμός («{exc}»).")
    if commission is None and commission_pct is not None:
        commission = round(gross * commission_pct / 100, 2)
    commission = commission or 0.0
    if gross < 0 or commission < 0:
        return failed("Αρνητικό ποσό.")
    try:
        rooms = parse_int(_cell(row, columns, "rooms")) or 1
        adults = parse_int(_cell(row, columns, "adults"))
        children = parse_int(_cell(row, columns, "children"))
        people = parse_int(_cell(row, columns, "people"))
    except _BadValue as exc:
        return failed(f"Τα δωμάτια και τα άτομα πρέπει να είναι ακέραιοι («{exc}»).")
    if adults is None and people is not None:
        adults = max(people - (children or 0), 0)
    channel = BOOKING_COM if fmt == "booking" else (_label(_cell(row, columns, "channel")) or DEFAULT_CHANNEL)[:100]

    booking = {
        "booking_ref": reference,
        "room_label": label,
        "channel": channel,
        "status": status,
        "booked_on": booked_on,
        "cancelled_on": cancelled_on if status == "cancelled" else None,
        "arrival": arrival,
        "departure": departure,
        "rooms": rooms,
        "adults": adults,
        "children": children,
        "gross_eur": gross,
        "commission_eur": commission,
        "country_code": country_code(_cell(row, columns, "country")),
    }
    return [ParsedRow(number, part["booking_ref"], part["room_label"], part) for part in _split_rooms(booking)]


def _split_rooms(booking: dict[str, Any]) -> list[dict[str, Any]]:
    """A Booking.com booking of several room types («Double Room, Twin Room», 2 rooms) → one booking per room.

    Only when the label lists exactly as many names as rooms; the same name
    repeated is one booking of that many rooms. Amounts and adults are
    shared evenly, the remainder on the last room. References get «-1», «-2».
    """
    rooms = booking["rooms"]
    parts = [part.strip() for part in booking["room_label"].split(",")]
    if rooms < 2 or len(parts) != rooms or not all(parts):
        return [booking]
    if len(set(parts)) == 1:
        return [{**booking, "room_label": parts[0]}]

    def share(total: float | int | None, index: int, cents: bool) -> float | int | None:
        if total is None:
            return None
        if cents:
            each = int(round(total * 100)) // rooms
            return (int(round(total * 100)) - each * (rooms - 1)) / 100 if index == rooms - 1 else each / 100
        each = total // rooms
        return total - each * (rooms - 1) if index == rooms - 1 else each

    return [
        {
            **booking,
            "booking_ref": f"{booking['booking_ref']}-{index + 1}",
            "room_label": part,
            "rooms": 1,
            "gross_eur": share(booking["gross_eur"], index, True),
            "commission_eur": share(booking["commission_eur"], index, True),
            "adults": share(booking["adults"], index, False),
            "children": share(booking["children"], index, False),
        }
        for index, part in enumerate(parts)
    ]


def _read_template_extras(sheets: dict[str, list[tuple]], parsed: ParsedFile) -> None:
    """«Δωμάτια» and «Κλειστές περίοδοι» of a template, when present."""
    extras = ((ROOMS_SHEET, ROOM_COLUMNS, "rooms"), (CLOSURES_SHEET, CLOSURE_COLUMNS, "closures"))
    for title, columns_spec, target in extras:
        rows = sheets.get(title)
        if not rows:
            continue
        fields = {column.key: (column.header,) for column in columns_spec}
        required = tuple(column.key for column in columns_spec if column.required)
        header_index, columns = _find_header(rows, fields, required)
        if not columns:
            continue
        for offset, row in enumerate(rows[header_index + 1:]):
            number = header_index + 2 + offset
            values = {key: _cell(row, columns, key) for key in columns}
            if all(_empty(value) for value in values.values()):
                continue
            try:
                item = _room_row(values) if target == "rooms" else _closure_row(values)
            except _BadValue as exc:
                parsed.sheet_errors.append(
                    {"row": number, "booking_ref": None, "message": f"Φύλλο «{title}»: {exc}"}
                )
                continue
            getattr(parsed, target).append(item)


def _room_row(values: dict[str, object]) -> dict[str, Any]:
    label = _label(values.get("room_label"))
    if label is None:
        raise _BadValue("λείπει ο τύπος δωματίου.")
    try:
        rooms = parse_int(values.get("rooms"))
        valid_from, valid_to = parse_date(values.get("valid_from")), parse_date(values.get("valid_to"))
    except _BadValue as exc:
        raise _BadValue(f"«{label}»: μη έγκυρη τιμή («{exc}»).") from exc
    if rooms is None:
        raise _BadValue(f"«{label}»: λείπει ο αριθμός δωματίων.")
    if valid_from and valid_to and valid_to < valid_from:
        raise _BadValue(f"«{label}»: το «Ισχύει έως» είναι πριν από το «Ισχύει από».")
    return {"room_label": label, "rooms": rooms, "valid_from": valid_from, "valid_to": valid_to}


def _closure_row(values: dict[str, object]) -> dict[str, Any]:
    try:
        date_from, date_to = parse_date(values.get("date_from")), parse_date(values.get("date_to"))
        rooms = parse_int(values.get("rooms"))
    except _BadValue as exc:
        raise _BadValue(f"μη έγκυρη τιμή («{exc}»).") from exc
    if date_from is None or date_to is None:
        raise _BadValue("λείπει το «Από» ή το «Έως».")
    if date_to < date_from:
        raise _BadValue("το «Έως» είναι πριν από το «Από».")
    if rooms == 0:
        rooms = None
    return {
        "date_from": date_from,
        "date_to": date_to,
        "room_label": _label(values.get("room_label")),
        "rooms": rooms,
        "reason": _label(values.get("reason")),
    }


def labels_of(parsed: ParsedFile) -> dict[str, int]:
    """Each room label of the file with its number of bookings, most bookings first."""
    counts: dict[str, int] = {}
    for booking in parsed.bookings:
        counts[booking["room_label"]] = counts.get(booking["room_label"], 0) + 1
    for room in parsed.rooms:
        counts.setdefault(room["room_label"], 0)
    return dict(sorted(counts.items(), key=lambda item: (-item[1], item[0])))

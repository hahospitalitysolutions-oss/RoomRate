"""«Απόδοση»: the RoomRate bookings template and GET /api/v1/performance/template.xlsx."""

import io
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.dependencies import (
    AccountContext,
    get_account_context,
    get_onboarding_repository,
    get_performance_repository,
)
from api.main import app
from api.services.performance_template import (
    BOOKING_COLUMNS,
    BOOKINGS_SHEET,
    CLOSURES_SHEET,
    EXAMPLE_PREFIX,
    INSTRUCTIONS_SHEET,
    ROOMS_SHEET,
    STATUSES,
    XLSX_MEDIA_TYPE,
    build_template,
)

ACCOUNT_ID = UUID("00000000-0000-0000-0000-000000000001")
PROPERTY_ID = UUID("00000000-0000-0000-0000-000000000456")
ROOM_TYPES = [
    {"id": UUID("00000000-0000-0000-0000-0000000000a1"), "label": "Double Room with Sea View", "rooms": 5},
    {"id": UUID("00000000-0000-0000-0000-0000000000a2"), "label": "Junior Suite", "rooms": None},
]


def workbook(content: bytes):
    return load_workbook(io.BytesIO(content))


def headers(sheet) -> list[str]:
    return [cell.value for cell in sheet[1] if cell.value]


def test_template_has_the_four_sheets_with_the_bookings_sheet_open():
    book = workbook(build_template(ROOM_TYPES))

    assert book.sheetnames == [BOOKINGS_SHEET, ROOMS_SHEET, CLOSURES_SHEET, INSTRUCTIONS_SHEET]
    assert book.active.title == BOOKINGS_SHEET


def test_bookings_sheet_marks_the_required_columns_and_offers_lists():
    sheet = workbook(build_template(ROOM_TYPES))[BOOKINGS_SHEET]

    assert headers(sheet) == [column.header + (" *" if column.required else "") for column in BOOKING_COLUMNS]
    assert "Αριθμός κράτησης *" in headers(sheet) and "Ημερομηνία κράτησης" in headers(sheet)
    rules = {rule.formula1: rule for rule in sheet.data_validations.dataValidation}
    status = rules['"' + ",".join(STATUSES) + '"']
    assert status.errorStyle in (None, "stop")  # a status outside the list is refused
    channel = next(rule for formula, rule in rules.items() if "Booking.com" in formula)
    assert channel.errorStyle == "warning"  # any other channel may stay
    assert f"'{ROOMS_SHEET}'!$A$2:$A$200" in rules
    assert sheet.freeze_panes == "A2"


def test_examples_are_grey_and_marked_for_the_importer_to_skip():
    sheet = workbook(build_template(ROOM_TYPES))[BOOKINGS_SHEET]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]

    assert len(rows) == 3
    assert all(row[0].startswith(EXAMPLE_PREFIX) for row in rows)
    assert {row[4] for row in rows} <= {room["label"] for room in ROOM_TYPES}
    assert sheet["A2"].font.italic and sheet["A2"].font.color.rgb.endswith("8A8F98")
    assert rows[2][11] == "Ακυρωμένη"  # the status is written as the owner reads it


def test_rooms_sheet_is_prefilled_with_the_property_room_types():
    sheet = workbook(build_template(ROOM_TYPES))[ROOMS_SHEET]

    assert [[cell.value for cell in row][:2] for row in sheet.iter_rows(min_row=2)] == [
        ["Double Room with Sea View", 5],
        ["Junior Suite", None],
    ]


def test_given_bookings_replace_the_examples_and_dates_are_dates():
    from datetime import date

    booking = {
        "booking_ref": "4011", "booked_on": date(2026, 1, 5), "arrival": date(2026, 7, 1),
        "departure": date(2026, 7, 4), "room_label": "Junior Suite", "gross_eur": 450.0,
        "channel": "Booking.com", "status": "no_show", "country": "Γερμανία",
    }
    sheet = workbook(build_template(ROOM_TYPES, bookings=[booking]))[BOOKINGS_SHEET]

    assert sheet.max_row == 2
    assert sheet["A2"].value == "4011" and sheet["L2"].value == "No-show"
    assert sheet["C2"].value.date() == date(2026, 7, 1) and sheet["C2"].number_format == "DD/MM/YYYY"


def test_instructions_explain_every_bookings_column_and_never_ask_for_guest_data():
    sheet = workbook(build_template(ROOM_TYPES, property_name="Aigaio Rooms"))[INSTRUCTIONS_SHEET]
    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)

    assert "Aigaio Rooms" in text
    for column in BOOKING_COLUMNS:
        assert column.header in text
    assert "δεν διαβάζονται ποτέ" in text


# ---------------------------------------------------------------------------
# GET /api/v1/performance/template.xlsx
# ---------------------------------------------------------------------------


class FakeOnboarding:
    def get_owned_property(self, account_id, owned_property_id):
        if account_id == ACCOUNT_ID and owned_property_id == PROPERTY_ID:
            return {"id": PROPERTY_ID, "display_name": "Aigaio Rooms"}
        return None


class FakePerformance:
    def __init__(self):
        self.calls = []

    def room_types(self, account_id, owned_property_id):
        self.calls.append((account_id, owned_property_id))
        return ROOM_TYPES


@pytest.fixture
def client():
    repository = FakePerformance()
    app.dependency_overrides[get_account_context] = lambda: AccountContext(account_id=ACCOUNT_ID)
    app.dependency_overrides[get_onboarding_repository] = lambda: FakeOnboarding()
    app.dependency_overrides[get_performance_repository] = lambda: repository
    yield TestClient(app), repository
    app.dependency_overrides.clear()


def test_template_endpoint_downloads_the_property_template(client):
    http, repository = client
    response = http.get("/api/v1/performance/template.xlsx", params={"owned_property_id": str(PROPERTY_ID)})

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="RoomRate_kratiseis.xlsx"'
    book = workbook(response.content)
    assert book[ROOMS_SHEET]["A2"].value == "Double Room with Sea View"
    assert repository.calls == [(ACCOUNT_ID, PROPERTY_ID)]


def test_template_endpoint_hides_another_account_property(client):
    http, repository = client
    response = http.get(
        "/api/v1/performance/template.xlsx",
        params={"owned_property_id": "00000000-0000-0000-0000-000000000999"},
    )

    assert response.status_code == 404
    assert response.json() == {"detail": "Owned property not found"}
    assert repository.calls == []


def test_calls_not_built_yet_answer_fastapi_not_found_for_the_page_samples(client):
    """The page falls back to public/mock/performance only on FastAPI's own 404 «Not Found»."""
    http, _ = client
    response = http.get("/api/v1/performance/not-built-yet", params={"owned_property_id": str(PROPERTY_ID)})

    assert response.status_code == 404
    assert response.json() == {"detail": "Not Found"}

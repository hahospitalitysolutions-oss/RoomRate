"""«Απόδοση»: synthetic bookings (api/services/performance_demo.py) and the files written from them."""

import csv
import importlib.util
import io
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from api.services.performance_countries import COUNTRIES
from api.services.performance_demo import (
    BOOKING_EXPORT_COLUMNS,
    CHANNELS,
    DEFAULT_ROOM_TYPES,
    booking_export_rows,
    generate_demo_bookings,
    template_rows,
)
from api.services.performance_template import BOOKINGS_SHEET, CLOSURES_SHEET, ROOMS_SHEET

AS_OF = date(2026, 10, 20)
SEASONS = (2025, 2026, 2027)
ROOMS = {room.label: room.rooms for room in DEFAULT_ROOM_TYPES}


@pytest.fixture(scope="module")
def demo():
    return generate_demo_bookings(seasons=SEASONS, as_of=AS_OF, seed=7)


def nights(booking):
    return [booking["arrival"] + timedelta(days=offset) for offset in range((booking["departure"] - booking["arrival"]).days)]


def test_same_seed_same_bookings():
    first = generate_demo_bookings(seasons=SEASONS, as_of=AS_OF, seed=3)
    again = generate_demo_bookings(seasons=SEASONS, as_of=AS_OF, seed=3)

    assert first.bookings == again.bookings
    assert generate_demo_bookings(seasons=SEASONS, as_of=AS_OF, seed=4).bookings != first.bookings


def test_no_room_is_sold_twice_for_a_night(demo):
    held = Counter(
        (booking["room_label"], night)
        for booking in demo.bookings
        if booking["status"] != "cancelled"
        for night in nights(booking)
    )

    assert held and all(count <= ROOMS[label] for (label, _), count in held.items())


def test_only_what_was_booked_by_the_day_exists(demo):
    assert max(booking["booked_on"] for booking in demo.bookings) <= AS_OF
    cancelled = [booking for booking in demo.bookings if booking["status"] == "cancelled"]
    assert cancelled
    assert all(booking["booked_on"] <= booking["cancelled_on"] <= AS_OF for booking in cancelled)
    assert all(booking["cancelled_on"] < booking["arrival"] for booking in cancelled)
    assert all(booking["arrival"] < AS_OF for booking in demo.bookings if booking["status"] == "no_show")


def test_stays_are_in_season_and_the_winters_are_closed(demo):
    assert all(5 <= night.month <= 10 for booking in demo.bookings for night in nights(booking))
    assert [(closure["date_from"], closure["date_to"]) for closure in demo.closures] == [
        (date(2024, 11, 1), date(2025, 4, 30)),
        (date(2025, 11, 1), date(2026, 4, 30)),
        (date(2026, 11, 1), date(2027, 4, 30)),
    ]


def test_high_summer_fills_and_costs_more_than_may_and_the_next_season_is_partly_booked(demo):
    sold, revenue = Counter(), Counter()
    for booking in demo.bookings:
        if booking["status"] != "confirmed":
            continue
        for night in nights(booking):
            sold[(night.year, night.month)] += 1
            revenue[(night.year, night.month)] += booking["gross_eur"] / len(nights(booking))

    assert sold[(2025, 8)] > 1.5 * sold[(2025, 5)]
    assert revenue[(2026, 8)] / sold[(2026, 8)] > revenue[(2026, 5)] / sold[(2026, 5)]
    assert revenue[(2026, 8)] / sold[(2026, 8)] > revenue[(2025, 8)] / sold[(2025, 8)]  # dearer than last year
    assert 0 < sold[(2027, 7)] < sold[(2026, 7)] / 2


def test_mix_of_channels_countries_and_cancellations(demo):
    channels = Counter(booking["channel"] for booking in demo.bookings)
    assert set(channels) == {name for name, _, _ in CHANNELS}
    assert channels.most_common(1)[0][0] == "Booking.com"
    assert all(booking["country_code"] in COUNTRIES for booking in demo.bookings)
    share = Counter(booking["status"] for booking in demo.bookings)["cancelled"] / len(demo.bookings)
    assert 0.05 < share < 0.15
    direct = [booking for booking in demo.bookings if booking["channel"] == "Απευθείας"]
    assert all(booking["commission_eur"] == 0 for booking in direct)
    assert len({booking["booking_ref"] for booking in demo.bookings}) == len(demo.bookings)


def test_template_rows_name_the_country_in_greek(demo):
    rows = template_rows(demo)

    assert {row["country"] for row in rows} <= {greek for greek, _ in COUNTRIES.values()}


def test_booking_export_lists_only_booking_com_as_the_extranet_does(demo):
    rows = booking_export_rows(demo)
    expected = [booking for booking in demo.bookings if booking["channel"] == "Booking.com"]

    assert len(rows) == len(expected)
    assert list(rows[0]) == list(BOOKING_EXPORT_COLUMNS)
    first, booking = rows[0], expected[0]
    assert first["Book number"] == booking["booking_ref"]
    assert first["Price"] == f"{booking['gross_eur']:.2f} EUR"
    assert first["Booker country"] == booking["country_code"].lower()
    assert {row["Status"] for row in rows} <= {"ok", "cancelled_by_guest", "no_show"}
    assert all(row["Guest name(s)"].startswith("Demo Guest") for row in rows)


# ---------------------------------------------------------------------------
# scripts/generate_demo_bookings.py
# ---------------------------------------------------------------------------


def _script():
    path = Path(__file__).resolve().parents[2] / "scripts" / "generate_demo_bookings.py"
    spec = importlib.util.spec_from_file_location("generate_demo_bookings", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_script_writes_the_template_and_the_booking_export(tmp_path, demo):
    paths = _script().write_files(demo, tmp_path)

    book = load_workbook(paths["template"])
    assert book[BOOKINGS_SHEET].max_row == len(demo.bookings) + 1
    assert book[ROOMS_SHEET]["B2"].value == DEFAULT_ROOM_TYPES[0].rooms
    assert book[CLOSURES_SHEET]["E2"].value == "Εκτός σεζόν"
    export = load_workbook(paths["export_xlsx"]).active
    assert [cell.value for cell in export[1]] == list(BOOKING_EXPORT_COLUMNS)
    with paths["export_csv"].open(encoding="utf-8-sig") as handle:
        assert len(list(csv.DictReader(handle))) == export.max_row - 1


def test_script_adds_three_broken_rows_on_request(tmp_path, demo):
    paths = _script().write_files(demo, tmp_path, with_errors=True)

    sheet = load_workbook(paths["template"])[BOOKINGS_SHEET]
    refs = [sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)]
    assert refs[-3:] == ["ERR-1", "ERR-2", "ERR-3"]
    text = paths["export_csv"].read_text(encoding="utf-8-sig")
    assert "τριακόσια" in text and len(list(csv.DictReader(io.StringIO(text)))) == len(booking_export_rows(demo)) + 3
